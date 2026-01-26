from models import ExperienceSection
from docx_lists import list_number
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Pt


import logging

def delete_paragraph(paragraph):
    """Helper function to remove a paragraph from a docx document."""
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None

def resolve_bullet_style(doc, style_type=WD_STYLE_TYPE.PARAGRAPH):
    """
    Resolves a bullet paragraph style that already exists in the template.
    If none exists, creates a basic style and signals that a text prefix is needed.
    """
    preferred = ["List Bullet", "List Bullet 2", "List Bullet 3", "List Paragraph"]
    for style_name in preferred:
        if style_name in doc.styles:
            return style_name, False

    for style in doc.styles:
        if style.type == style_type and "bullet" in style.name.lower():
            return style.name, False

    style_name = "List Bullet"
    try:
        logging.info(f"Style '{style_name}' not found. Creating it.")
        base_style = doc.styles["Normal"]
        new_style = doc.styles.add_style(style_name, style_type)
        new_style.base_style = base_style

        # Match common list paragraph spacing/indentation.
        new_style.font.name = "Calibri"
        new_style.font.size = Pt(10)
        p_fmt = new_style.paragraph_format
        p_fmt.left_indent = Pt(36)
        p_fmt.first_line_indent = Pt(-18)
        p_fmt.space_after = Pt(0)
        p_fmt.line_spacing = 1.0

        logging.info(f"Style '{style_name}' created successfully.")
        return style_name, True
    except Exception as e:
        logging.error(f"Error adding style '{style_name}': {e}")
        fallback = "List Paragraph" if "List Paragraph" in doc.styles else "Normal"
        return fallback, True

def replace_section(doc, heading_text: str, new_content: ExperienceSection):
    """
    Replaces a section in the document with new, formatted content.
    """
    bullet_style, needs_prefix = resolve_bullet_style(doc)
    needs_prefix = False
    start_idx = -1
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip().upper() == heading_text.upper():
            start_idx = i
            break
    
    if start_idx == -1:
        logging.warning(f"Section heading '{heading_text}' not found.")
        return

    end_idx = -1
    for i in range(start_idx + 1, len(doc.paragraphs)):
        if doc.paragraphs[i].text.strip().isupper():
            end_idx = i
            break
    
    if end_idx == -1:
        end_idx = len(doc.paragraphs)

    for i in range(end_idx - 1, start_idx, -1):
        delete_paragraph(doc.paragraphs[i])

    anchor = doc.paragraphs[start_idx + 1] if start_idx + 1 < len(doc.paragraphs) else None

    for exp in new_content.experiences:
        # Add Company and Duration
        if anchor:
            p_company_duration = anchor.insert_paragraph_before()
        else:
            p_company_duration = doc.add_paragraph()

        p_fmt = p_company_duration.paragraph_format
        p_fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
        right_pos = (
            doc.sections[0].page_width
            - doc.sections[0].left_margin
            - doc.sections[0].right_margin
        )
        p_fmt.tab_stops.add_tab_stop(
            right_pos,
            WD_TAB_ALIGNMENT.RIGHT,
            WD_TAB_LEADER.SPACES,
        )

        company_run = p_company_duration.add_run(exp.company)
        company_run.bold = True
        p_company_duration.add_run(f"\t{exp.duration}")

        # Add Role
        if anchor:
            p_role = anchor.insert_paragraph_before()
        else:
            p_role = doc.add_paragraph()
        role_run = p_role.add_run(exp.role)
        role_run.bold = True
        role_run.underline = True

        # Add bullets
        if exp.bullets:
            prev_bullet = None
            for bullet in exp.bullets:
                if anchor:
                    p_bullet = anchor.insert_paragraph_before("", style=bullet_style)
                else:
                    p_bullet = doc.add_paragraph("", style=bullet_style)
                if needs_prefix:
                    p_bullet.add_run("- ")
                else:
                    list_number(doc, p_bullet, prev=prev_bullet, level=0, num=False)
                for segment in bullet:
                    run = p_bullet.add_run(segment.text)
                    if segment.bold:
                        run.bold = True
                prev_bullet = p_bullet
