import unittest
import docx
import sys
import os
import logging

# Configure logging to stdout.
logging.basicConfig(stream=sys.stdout, level=logging.INFO)

# Add the project root directory to the Python path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from models import Experience, ExperienceSection, TextSegment
from resume import replace_section

class TestReplaceSection(unittest.TestCase):
    def _find_paragraph_index(self, paragraphs, text):
        for idx, paragraph in enumerate(paragraphs):
            if paragraph.text == text:
                return idx
        return -1

    def test_replace_experience_section(self):
        """
        Tests the replace_section logic with a mock ExperienceSection.
        """
        # 1. Create mock data with 4 experiences and mixed bold/normal text
        mock_data = ExperienceSection(
            experiences=[
                Experience(
                    company="Tech Innovations Inc.",
                    role="Lead Software Developer",
                    duration="2020 - Present",
                    bullets=[
                        [TextSegment(text="Architected and implemented a new "), TextSegment(text="microservices-based platform", bold=True), TextSegment(text=".")],
                        [TextSegment(text="Improved API response times by "), TextSegment(text="40%", bold=True), TextSegment(text=" through query optimization.")],
                    ],
                ),
                Experience(
                    company="Data Driven Co.",
                    role="Senior Data Analyst",
                    duration="2018 - 2020",
                    bullets=[
                        [TextSegment(text="Developed ETL pipelines that processed over "), TextSegment(text="1TB of data daily", bold=True), TextSegment(text=".")],
                        [TextSegment(text="Built dashboards for key business metrics.")],
                        [TextSegment(text="Created and maintained data models.")],
                        [TextSegment(text="Performed statistical analysis on large datasets.")],
                        [TextSegment(text="Presented findings to stakeholders.")],
                    ],
                ),
                Experience(
                    company="Web Solutions LLC",
                    role="Frontend Developer",
                    duration="2016 - 2018",
                    bullets=[
                        [TextSegment(text="Created responsive user interfaces using "), TextSegment(text="React and Redux", bold=True), TextSegment(text=".")],
                        [TextSegment(text="Worked with designers to create a consistent look and feel.")],
                        [TextSegment(text="Improved website performance by optimizing assets.")],
                        [TextSegment(text="Wrote unit and integration tests.")],
                        [TextSegment(text="Participated in code reviews.")],
                    ],
                ),
                Experience(
                    company="Startup Seed",
                    role="Junior Developer",
                    duration="2015 - 2016",
                    bullets=[
                        [TextSegment(text="Fixed bugs and implemented minor features.")],
                        [TextSegment(text="Gained experience with "), TextSegment(text="Agile methodologies", bold=True), TextSegment(text=".")],
                    ],
                ),
            ]
        )

        # 2. Create a dummy docx document
        doc = docx.Document()
        doc.add_paragraph("John Doe", style='Title')
        doc.add_paragraph("WORK EXPERIENCE", style='Heading 1')
        doc.add_paragraph("This is old content that should be deleted.")
        doc.add_paragraph("Another line of old content.")
        doc.add_paragraph("SKILLS", style='Heading 1')
        doc.add_paragraph("Python, SQL, Docker")

        # 3. Call the function to be tested
        replace_section(doc, "WORK EXPERIENCE", mock_data)

        # 4. Assert the document content is correct
        paragraphs = doc.paragraphs
        paragraph_texts = [p.text for p in paragraphs]

        # Check that old content is gone and new content is present
        self.assertNotIn("This is old content that should be deleted.", paragraph_texts)
        self.assertIn("Tech Innovations Inc.\t2020 - Present", paragraph_texts)
        self.assertIn("Lead Software Developer", paragraph_texts)
        self.assertIn("Architected and implemented a new microservices-based platform.", paragraph_texts)
        self.assertIn("Data Driven Co.\t2018 - 2020", paragraph_texts)
        self.assertIn("Senior Data Analyst", paragraph_texts)
        self.assertIn("Developed ETL pipelines that processed over 1TB of data daily.", paragraph_texts)
        self.assertIn("Web Solutions LLC\t2016 - 2018", paragraph_texts)
        self.assertIn("Frontend Developer", paragraph_texts)
        self.assertIn("Created responsive user interfaces using React and Redux.", paragraph_texts)
        self.assertIn("Startup Seed\t2015 - 2016", paragraph_texts)

        # Assert ordering within an experience block: company/duration -> role -> bullets
        company_idx = self._find_paragraph_index(paragraphs, "Tech Innovations Inc.\t2020 - Present")
        role_idx = self._find_paragraph_index(paragraphs, "Lead Software Developer")
        first_bullet_idx = self._find_paragraph_index(
            paragraphs,
            "Architected and implemented a new microservices-based platform.",
        )
        self.assertNotEqual(company_idx, -1)
        self.assertNotEqual(role_idx, -1)
        self.assertNotEqual(first_bullet_idx, -1)
        self.assertLess(company_idx, role_idx)
        self.assertLess(role_idx, first_bullet_idx)

        # Deeper inspection of formatting (first experience)
        company_para_found = False
        for para in paragraphs:
            if para.text == "Tech Innovations Inc.\t2020 - Present":
                self.assertTrue(para.runs[0].bold)
                company_para_found = True
            if para.text == "Lead Software Developer":
                self.assertTrue(para.runs[0].bold)
                self.assertTrue(para.runs[0].underline)
            if para.text == "Architected and implemented a new microservices-based platform.":
                # Expected: "microservices-based platform" is bold
                self.assertEqual(len(para.runs), 3)
                self.assertFalse(para.runs[0].bold)
                self.assertTrue(para.runs[1].bold)
                self.assertEqual(para.runs[1].text, "microservices-based platform")
                self.assertFalse(para.runs[2].bold)

        self.assertTrue(company_para_found, "Company paragraph not found for detailed check.")

if __name__ == '__main__':
    unittest.main()
