from fastapi import FastAPI, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
from typing import Optional
from bs4 import BeautifulSoup
from database import create_resumes_table, insert_resume
from engine import generate_with_gemini
from resume import replace_section

import docx
from io import BytesIO
import logging
import requests


app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173", "http://192.168.1.67:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class Resume(BaseModel):
    file_path: str


def setup_logging():
    """Sets up logging to a file."""
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=[
            logging.FileHandler("logs/app.log"),
            logging.StreamHandler(),
        ],
    )

setup_logging()

def scrape_job_description(url: str) -> str:
    try:
        response = requests.get(url)
        response.raise_for_status()  # Raise an exception for bad status codes
        soup = BeautifulSoup(response.content, "html.parser")
        # This is a generic way to get text. It might need to be adjusted for specific websites.
        # We can try to find common tags for job descriptions like 'div' with id 'job-description' or class 'job-description'
        job_description = soup.find("div", id="job-description")
        if not job_description:
            job_description = soup.find("div", class_="job-description")

        if job_description:
            return job_description.get_text()
        else:
            # As a fallback, get all the text from the body
            if soup.body:
                return soup.body.get_text()
            else:
                logging.error("No body tag found in the HTML.")
                return ""
    except requests.exceptions.RequestException as e:
        logging.error(f"Error scraping job description: {e}")
        return ""

def extract_docx_text(doc):
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    return '\n'.join(full_text)


def sanitize_filename(url: str) -> str:
    """
    Sanitizes a URL to create a valid filename.
    """
    # Get the last part of the URL
    filename = url.split('/')[-1]
    # Remove invalid characters
    invalid_chars = ['?', '=', '&', ':', '/', '\\']
    for char in invalid_chars:
        filename = filename.replace(char, '_')
    return filename

@app.on_event("startup")
def on_startup():
    create_resumes_table()

def build_resume(job_url: str, template_bytes: Optional[bytes]):
    # 1. Scrape job description from job.url
    job_description = scrape_job_description(job_url)
    if not job_description:
        return None, "Could not scrape job description"

    # 2. Load the source resume template
    try:
        if template_bytes:
            doc = docx.Document(BytesIO(template_bytes))
            resume_text = extract_docx_text(doc)
        else:
            source_resume_path = "/app/resume_templates/source_resume.docx"
            doc = docx.Document(source_resume_path)
            resume_text = extract_docx_text(doc)
    except FileNotFoundError:
        return None, "Source resume not found"
    except Exception:
        return None, "Invalid or unreadable resume template"

    # 3. Use Gemini to generate ATS-beating experience and skills
    new_experience = generate_with_gemini(job_description, resume_text)

    # 4. Modify the resume
    replace_section(doc, "WORK EXPERIENCE", new_experience)

    # 5. Save the new resume
    file_name = f"resume_for_{sanitize_filename(job_url)}.docx"
    file_path = f"generated_resumes/{file_name}"
    doc.save(file_path)

    # 6. Store information in the database
    insert_resume(job_url, file_path)
    return file_path, None


@app.post("/generate-resume/", response_model=Resume)
async def generate_resume(url: str = Form(...), template: Optional[UploadFile] = File(default=None)):
    template_bytes = await template.read() if template else None
    file_path, error = build_resume(url, template_bytes)
    if error:
        return {"error": error}
    if not file_path:
        return {"error": "Failed to generate resume"}
    return Resume(file_path=file_path)


@app.post("/generate-resume-file/")
async def generate_resume_file(url: str = Form(...), template: Optional[UploadFile] = File(default=None)):
    template_bytes = await template.read() if template else None
    file_path, error = build_resume(url, template_bytes)
    if error:
        return {"error": error}
    if not file_path:
        return {"error": "Failed to generate resume"}
    filename = file_path.split("/")[-1]
    return FileResponse(
        path=file_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=filename,
    )

@app.get("/")
async def root():
    return {"message": "Welcome to the Resumer API!"}
